#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""f3_fill_policy 的回归测试 —— 钉住「空格从哪来、怎么消灭」这件事。

★ 为什么这些用例必须逐值断言（不能只查不变量）：
    本次的 bug 恰恰是**所有不变量都满足**的情况下发生的 ——
    阴性/待确认的 `值` 本来就是 None，写空串是"合法"的，
    覆盖率指标也完全看不出来。只有逐格断言具体值才抓得住。
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import f3_fill_policy as fp                                  # noqa: E402

PASS = FAIL = 0


def ck(name, got, want):
    global PASS, FAIL
    if got == want:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name}\n      得到 {got!r}\n      期望 {want!r}")


def ok(name, cond):
    ck(name, bool(cond), True)


print("── 1. render：按 state 渲染（本次修复的核心）──")
G = "词条出现但被否定（前置或后置）"          # route_f3.py:594 的有据阴性
ck("阴性+check(有据)", fp.render("阴性", None, "check", why=G), "未做")
ck("阴性+lab(有据)", fp.render("阴性", None, "lab", why=G), "未查")
ck("阴性+drug(有据)", fp.render("阴性", None, "drug", why=G), "未用")
ck("阴性+bool(有据)", fp.render("阴性", None, "bool", why=G), "否")
ck("阴性+conditional(有据)", fp.render("阴性", None, "conditional", why=G), "不适用")
ck("阴性+manual(有据)", fp.render("阴性", None, "manual", why=G), "待人工")
ck("阴性+text(有据)", fp.render("阴性", None, "text", why=G), "未记录")
ck("阴性+未知kind 兜底", fp.render("阴性", None, "乱写", why=G),
   fp.NEG_TEXT[fp.DEFAULT_KIND])

print("── 1b. 阴性分两路（route_f3.py:1224「阴性不是一个，是两个」）──")
# ★ 这是被纠出来的真 bug：把配置断言渲染成「未做」= 把猜测当事实。
#   56 个零命中列的整片阴性全是配置断言 —— 100 个卒中病人「头颅CT 未做」
#   在临床上不可能是真的。
CFG = "来源段存在（辅助检查）却没命中 → 否"     # route_f3.py:1029
ck("配置断言+check → 未提及（不断言没做）", fp.render("阴性", None, "check", why=CFG),
   "未提及")
ck("配置断言+bool → 未提及", fp.render("阴性", None, "bool",
   why="配置断言「缺失即否」"), "未提及")
ck("体健式兜底 → 未提及", fp.render("阴性", None, "bool",
   why="未命中任何阳性证据，且 既往史 为体健式表述"), "未提及")
ck("不适用 why → 不适用", fp.render("阴性", None, "conditional",
   why="不适用（条件列/备注列非必填），无内容即留空"), "不适用")
ck("**无 why 不敢断言** → 未提及", fp.render("阴性", None, "check"), "未提及")
ck("乱写的 why 也不敢断言", fp.render("阴性", None, "check", why="嗯"), "未提及")

print("── 1c. neg_grounded / neg_is_na ──")
ok("有据识别", fp.neg_grounded(G))
for w in (CFG, "配置断言「缺失即否」", "未命中任何阳性证据，且 x 为体健式表述",
          "", None, "嗯"):
    ok(f"没据识别：{w!r:.30}", not fp.neg_grounded(w))
ok("不适用识别", fp.neg_is_na("不适用（条件列/备注列非必填），无内容即留空"))
ok("有据不是不适用", not fp.neg_is_na(G))

print("── 2. 待确认绝不写空（项目铁律：静默失败是最坏的错）──")
ck("待确认", fp.render("待确认", None, "bool"), "待确认")
ck("state 缺失也算待确认", fp.render(None, None, "bool"), "待确认")
ck("未知 state 也算待确认", fp.render("乱", None, "bool"), "待确认")
for st in ("阴性", "待确认", None, "乱"):
    ok(f"render({st!r}) 非空", fp.render(st, None, "bool") != "")

print("── 3. 阳性：值照写，占位词保留（它是答案，不是空）──")
ck("阳性有值", fp.render("阳性", "2024-09-14", "text"), "2024-09-14")
ck("阳性空值→未提及", fp.render("阳性", None, "text"), "未提及")
ck("阳性占位词保留", fp.render("阳性", "未提及", "text"), "未提及")

print("── 4. compat 模式 = 旧行为（可回滚 / 可 A/B）──")
ck("compat 阴性→空", fp.render("阴性", None, "bool", "compat"), "")
ck("compat 待确认→空", fp.render("待确认", None, "bool", "compat"), "")
ck("compat 阳性→值", fp.render("阳性", "张三", "text", "compat"), "张三")

print("── 5. col_kind：列类型分派（实测踩过的那几个）──")
# 实测踩过：「抗血小板药物史」被 lab 规则里的「血小板」命中，误判成检验列。
# 它是既往史布尔列（以「史」结尾），正确归属是 bool —— 关键是**不能是 lab**。
ck("抗血小板药物史→bool（不是 lab）", fp.col_kind("抗血小板药物史"), "bool")
ok("抗血小板药物史 不是 lab", fp.col_kind("抗血小板药物史") != "lab")
ck("既往高同型半胱氨酸血症→bool（不是 lab）",
   fp.col_kind("既往高同型半胱氨酸血症"), "bool")
ck("病案号→text（不是 bool）", fp.col_kind("病案号"), "text")
ck("姓名→text", fp.col_kind("姓名"), "text")
ck("既往高血压→bool", fp.col_kind("既往高血压"), "bool")
ck("是否存在×禁忌症→bool", fp.col_kind("是否存在抗血小板治疗禁忌症"), "bool")
ck("头颅DWI检查→check", fp.col_kind("头颅弥散成像（DWI)检查"), "check")
ck("颈动脉B超→check", fp.col_kind("颈动脉B超检查"), "check")
ck("空腹血糖→lab", fp.col_kind("空腹血糖（mg/dL）"), "lab")
ck("凝血酶原时间→lab", fp.col_kind("凝血酶原时间PT（s，>50）"), "lab")
ck("住院阿司匹林用药时间→drug", fp.col_kind("住院阿司匹林用药时间"), "drug")
ck("住院天数不是药→text", fp.col_kind("住院天数（天）"), "text")
ck("住院总费用不是药→text", fp.col_kind("住院总费用（元）"), "text")
ck("序号→manual", fp.col_kind("序号"), "manual")
ck("质控打分→manual", fp.col_kind("质控打分"), "manual")
ck("仅TIA→conditional", fp.col_kind("症状持续时间(仅TIA需要填写)"), "conditional")
ck("醒后卒中→conditional", fp.col_kind("最后看起来正常时间(适用于醒来卒中)"), "conditional")
ck("空列名不炸", fp.col_kind(""), fp.DEFAULT_KIND)

print("── 6. 兜底类型是 text，不是 bool（否则「病案号」会被写成「否」）──")
ck("DEFAULT_KIND", fp.DEFAULT_KIND, "text")
ok("兜底渲染不是「否」", fp.render("阴性", None, fp.DEFAULT_KIND) != "否")

print("── 7. 占位词识别 ──")
for w in ("未提及", "无", "unknown", "—"):
    ok(f"{w!r} 是占位", fp.is_placeholder(w))
for w in ("2024-09-14", "否", "未做", "3.7", ""):
    ok(f"{w!r} 不是占位", not fp.is_placeholder(w))

print("── 8. fill_stats：两个口径必须分开 ──")
s = fp.fill_stats(["阳性", "阳性", "阴性", "待确认"],
                  ["张三", "未提及", None, None])
ck("总格数", s["总格数"], 4)
ck("阳性", s["阳性"], 2)
ck("占位", s["占位"], 1)
ck("阴性", s["阴性"], 1)
ck("待确认", s["待确认"], 1)
ck("有效填充率 = (2-1)/4", round(s["有效填充率"], 4), 0.25)
ck("覆盖率 = (2+1)/4", round(s["覆盖率"], 4), 0.75)
ok("覆盖率 > 有效填充率（差在阴性+占位）", s["覆盖率"] > s["有效填充率"])

print("── 8b. 有据覆盖率：配置断言的阴性不算结论 ──")
s2 = fp.fill_stats(
    ["阴性", "阴性", "阳性"],
    ["未提及", "否", "张三"],
    ["来源段存在（辅助检查）却没命中 → 否",       # 配置断言
     "词条出现但被否定（前置或后置）",             # 有据
     None])
ck("阴性_有据", s2["阴性_有据"], 1)
ck("阴性_配置断言", s2["阴性_配置断言"], 1)
ck("有据覆盖率 = (1-0+1)/3", round(s2["有据覆盖率"], 4), 0.6667)
ok("有据覆盖率 ≤ 覆盖率", s2["有据覆盖率"] <= s2["覆盖率"])
# 全是配置断言时，有据覆盖率应该塌到只剩阳性 —— 这正是 56 个零命中列的样子
s3 = fp.fill_stats(["阴性"] * 4, [None] * 4,
                   ["配置断言「缺失即否」"] * 4)
ck("全配置断言 → 有据覆盖率 0", round(s3["有据覆盖率"], 4), 0.0)
ck("全配置断言 → 覆盖率仍 100%（**虚的**）", round(s3["覆盖率"], 4), 1.0)
# 不传 whys 时退化为旧口径（兼容）
s4 = fp.fill_stats(["阴性", "阳性"], [None, "x"])
ck("不传 whys → 有据阴性按 0 算", s4["阴性_有据"], 0)

print("── 9. 门禁 check_gates ──")
good = {"总格数": 100, "阳性": 30, "阴性": 60, "待确认": 10, "占位": 0,
        "有效填充率": 0.30, "覆盖率": 0.90}
ok("健康输入→通过", fp.check_gates(good, 10, [], [], [], 0)[0])

_, r = fp.check_gates(good, 10, [], [], ["a", "b"], 0)
ok("缺 2 例→拦下", any("没有对应输出行" in x for x in r))

bad = {"总格数": 100, "阳性": 5, "阴性": 5, "待确认": 90, "占位": 0,
       "有效填充率": 0.05, "覆盖率": 0.10}
_, r = fp.check_gates(bad, 10, [], [], [], 0)
ok("填充率过低→拦下", any("有效填充率" in x for x in r))
ok("覆盖率过低→拦下", any("覆盖率" in x for x in r))

_, r = fp.check_gates(good, 10, [f"c{i}" for i in range(90)], [], [], 0)
ok("零命中列过多→拦下", any("零命中列" in x for x in r))

_, r = fp.check_gates(good, 10, [], [("x",)] * 9, [], 0)
ok("高空值行过多→拦下", any("OCR 挂了" in x for x in r))

_, r = fp.check_gates(good, 10, [], [], [], 7)
ok("filled 模式有空格→拦下", any("空单元格" in x for x in r))

ok("空表→拦下", not fp.check_gates({"总格数": 0}, 0, [], [], [], 0)[0])

print("── 10. 端到端：106 源 → 104 行必须被拦下并点名 ──")
with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    flat = ["病案号", "既往高血压", "头颅CT检查", "住院阿司匹林用药时间", "序号"]
    cols = td / "cols.json"
    cols.write_text(__import__("json").dumps({"flat": flat}), encoding="utf-8")

    indir = td / "out"
    indir.mkdir()
    import json
    G = "词条出现但被否定（前置或后置）"
    CFG = "来源段存在（辅助检查）却没命中 → 否"
    for i in range(1, 4):                    # 只出 3 例
        # case2 的 CT 用配置断言 why —— 验证 证据.why 真的流进了渲染
        ct_why = CFG if i == 2 else G
        (indir / f"case{i}.json").write_text(json.dumps({
            "case": f"case{i}",
            "items": {"病案号": {"state": "阳性", "值": f"10{i}"},
                      "既往高血压": {"state": "阴性", "值": None,
                                     "证据": {"why": G}},
                      "头颅CT检查": {"state": "阴性", "值": None,
                                      "证据": {"why": ct_why}},
                      "住院阿司匹林用药时间": {"state": "待确认", "值": None},
                      "序号": {"state": "阴性", "值": None, "证据": {"why": G}}},
        }, ensure_ascii=False), encoding="utf-8")

    mani = td / "input_manifest.tsv"         # 源端 5 例
    mani.write_text("name\tdir\tfile\n"
                    + "".join(f"case{i}\tH/case{i}\tcase{i}.pdf\n" for i in range(1, 6)),
                    encoding="utf-8")

    sys.argv = ["x", "--indir", str(indir), "--out", str(td / "w.xlsx"),
                "--columns", str(cols), "--manifest", str(mani), "--no-trace"]
    import importlib
    wide = importlib.import_module("f3_to_wide")
    try:
        wide.main()
        ok("缺例时应 exit 2", False)
    except SystemExit as e:
        ck("缺例时退出码", e.code, 2)

    # --no-gate 强制写盘，并检查异常清单点名了缺的那 2 例
    sys.argv += ["--no-gate"]
    wide.main()
    from openpyxl import load_workbook
    wb = load_workbook(td / "w.xlsx")
    ok("有异常清单 sheet", "异常清单" in wb.sheetnames)
    ok("有取值说明 sheet", "取值说明" in wb.sheetnames)
    txt = "\n".join(str(c.value) for row in wb["异常清单"].iter_rows()
                    for c in row if c.value is not None)
    ok("点名缺了 case4", "case4" in txt)
    ok("点名缺了 case5", "case5" in txt)
    ok("差额是 2", "2" in txt)

    ws = wb["Sheet1"]
    hdr = [c.value for c in ws[1]]
    r1 = [c.value for c in ws[2]]        # case1（全有据）
    r2 = [c.value for c in ws[3]]        # case2（CT 是配置断言）
    # R1（2026-10-07）在 498 列**之后**追加 3 列病史来源列。
    # 前缀 5+len(flat) 必须一字不差 —— 追加只能发生在尾巴上。
    ck("表头列数", len(hdr), 5 + len(flat) + len(wide.HISTORY_SRC_NAMES))
    ck("既有列顺序未变", hdr[:5 + len(flat)],
       ["住院号", "姓名", "来源目录", "输入文件", "NIHSS评分"] + flat)
    ck("R1 来源列追加在末尾", hdr[5 + len(flat):], list(wide.HISTORY_SRC_NAMES))
    ck("既往高血压阴性→否", r1[hdr.index("既往高血压")], "否")
    ck("头颅CT检查阴性→未做", r1[hdr.index("头颅CT检查")], "未做")
    ck("用药时间待确认→待确认", r1[hdr.index("住院阿司匹林用药时间")], "待确认")
    ck("序号阴性→待人工", r1[hdr.index("序号")], "待人工")
    ok("整行无空串", all(str(v).strip() != "" for v in r1[5:]))
    # ★ 证据.why 真的流进了渲染：同一列，有据→未做、配置断言→未提及
    ck("case2 同列配置断言→未提及", r2[hdr.index("头颅CT检查")], "未提及")
    ok("case1≠case2（同列两路渲染）",
       r1[hdr.index("头颅CT检查")] != r2[hdr.index("头颅CT检查")])

    # 异常清单 ③ 要把「配置断言占比」写出来
    t3 = "\n".join(str(c.value) for row in wb["异常清单"].iter_rows()
                   for c in row if c.value is not None)
    ok("零命中列表里有「配置断言」列", "其中配置断言" in t3)

print()
print(f"{'='*46}\n  {PASS} 项通过 / {FAIL} 项失败\n{'='*46}")
sys.exit(1 if FAIL else 0)
